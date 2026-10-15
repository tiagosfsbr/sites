from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'Briefing de Informacoes para Atualizacao do Site ECamiza Contabilidade.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)

styles = doc.styles
styles['Normal'].font.name = 'Arial'; styles['Normal'].font.size = Pt(10.5)
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
for name, size in [('Title', 23), ('Heading 1', 15), ('Heading 2', 11.5)]:
    s = styles[name]; s.font.name = 'Arial'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(0,0,0)
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def set_cell_border(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        tag = 'w:' + edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:color'), color)
def add_field(label, lines=1, note=None):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label); r.bold = True
    if note: p.add_run('  ' + note).italic = True
    for _ in range(lines):
        q = doc.add_paragraph('________________________________________________________________________________')
        q.paragraph_format.space_after = Pt(4); q.paragraph_format.line_spacing = 1.0
def add_choice(label, options):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    p.add_run(label + ' ').bold = True
    p.add_run('    '.join('☐ ' + x for x in options))
def heading(text, level=1):
    p = doc.add_paragraph(text, style='Heading ' + str(level)); p.paragraph_format.space_before = Pt(14 if level == 1 else 7); p.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run('Formulario de Informacoes para Atualizacao do Site')
sub = doc.add_paragraph('ECamiza Contabilidade'); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.runs[0].bold = True; sub.runs[0].font.size = Pt(13)
p = doc.add_paragraph('Preencha o que já estiver definido. Se algum item ainda não existir, escreva “não definido”. Envie também os arquivos e links indicados ao final.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(14)

heading('1 Dados da empresa')
add_field('Nome de divulgação no site')
add_field('Razão social')
add_field('CNPJ')
add_field('Número do CRC e estado')
add_field('Nome e cargo da pessoa responsável pela aprovação do site')
add_field('E-mail da pessoa responsável')

heading('2 Contatos e atendimento')
add_field('WhatsApp comercial com DDD e país', 1, 'Exemplo: +55 12 99999 9999')
add_field('Telefone comercial com DDD')
add_field('E-mail público de atendimento')
add_field('Horário de atendimento')
add_field('Instagram', 1, 'Informe o @ e o link completo')
add_field('Outras redes ou links relevantes', 2)
add_field('Mensagem desejada para o primeiro contato no WhatsApp', 2)

heading('3 Endereço e área de atendimento')
add_field('Endereço completo para exibição no site')
add_field('Link da ficha da empresa no Google Maps')
add_field('Link para traçar rota no Google Maps', 1, 'Se houver')
add_choice('Atendimento', ['Somente presencial', 'Somente online', 'Presencial e online'])
add_field('Cidades estados ou regiões atendidas', 2)

heading('4 Posicionamento e apresentação')
add_field('Em uma frase, como deseja que a empresa seja apresentada', 2)
add_field('Quais são os três principais diferenciais do escritório', 3)
add_field('Tom de comunicação desejado', 1, 'Exemplo: acolhedor, técnico, sofisticado, direto')
add_field('Público principal e perfil de cliente ideal', 3)
add_field('Segmentos que atende com prioridade', 2)
add_field('Segmentos ou serviços que não atende', 2)

heading('5 Serviços')
doc.add_paragraph('Marque os serviços que devem aparecer e descreva qualquer observação importante.')
services = ['Abertura e regularização de empresas', 'Contabilidade mensal', 'Planejamento tributário', 'Departamento pessoal e folha', 'Fiscal e obrigações acessórias', 'Societário e alterações contratuais', 'Imposto de renda pessoa física', 'Reforma tributária', 'Consultoria financeira ou gestão', 'Outros']
for s in services:
    p = doc.add_paragraph('☐ ' + s + '  Observações: _________________________________________________')
    p.paragraph_format.space_after = Pt(3)
add_field('Serviços que merecem destaque na página inicial', 2)
add_field('Serviços, promessas ou termos que não podem aparecer', 2)

heading('6 Planos e propostas')
add_choice('Deseja mostrar planos no site', ['Sim', 'Não', 'Ainda não sei'])
add_field('Nome dos planos e o que inclui cada um', 4)
add_choice('Deseja exibir preços', ['Sim', 'Não', 'Somente a partir de'])
add_field('Condições ou avisos que devem acompanhar propostas e preços', 2)

heading('7 Materiais visuais')
add_field('Logo final', 1, 'Envie PNG com fundo transparente e, se possível, SVG ou PDF vetorial')
add_field('Cores ou manual de marca', 1, 'Envie códigos de cor ou o material da marca')
add_field('Fotos que podem ser usadas', 1, 'Envie fotos da Eliane, equipe, escritório e atendimento com autorização')
add_field('Vídeos que podem ser usados', 1, 'Informe se há vídeo institucional ou se prefere não usar vídeos')
add_field('Há alguma foto, pessoa ou imagem que não deve aparecer', 2)

heading('8 Conteúdo e prova social')
add_field('Breve história do escritório ou da profissional', 4)
add_field('Formação, especialidades, certificações e experiências relevantes', 3)
add_field('Números que podem ser divulgados', 1, 'Exemplo: anos de experiência, clientes atendidos. Use apenas dados reais.')
add_field('Depoimentos reais autorizados', 1, 'Enviar nome, empresa ou cargo, texto e autorização de uso')
add_field('Link de avaliações do Google', 1, 'Se desejar mostrar avaliações')

heading('9 Site e publicação')
add_field('Domínio desejado', 1, 'Exemplo: www.suaempresa.com.br')
add_field('A empresa já possui domínio', 1, 'Informe qual e quem possui acesso')
add_field('E-mail para receber formulários de contato', 1, 'Se for diferente do e-mail público')
add_choice('Calculadora tributária no site', ['Manter', 'Remover', 'Avaliar depois'])
add_choice('Política de privacidade', ['Já possui', 'Precisa criar', 'Não sei'])
add_field('Outras observações, referências de sites e pedidos específicos', 4)

heading('10 Checklist de envio')
for item in ['Logo final em boa qualidade', 'Dados de contato confirmados', 'CNPJ e CRC', 'Endereço e link do Google Maps', 'Fotos autorizadas', 'Links de redes sociais', 'Textos ou informações institucionais', 'Depoimentos reais autorizados', 'Domínio e acessos, se já existirem']:
    p = doc.add_paragraph('☐ ' + item); p.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph('Ao enviar este formulário, confirme que as informações, fotos, depoimentos e marcas fornecidos podem ser utilizados na criação e publicação do site.'); p.paragraph_format.space_before = Pt(12)
add_field('Nome da pessoa que preencheu')
add_field('Data')

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Briefing de atualização do site  |  ECamiza Contabilidade')
doc.core_properties.title = 'Formulario de Informacoes para Atualizacao do Site ECamiza Contabilidade'
doc.core_properties.author = 'Tiago'
doc.save(OUT)
